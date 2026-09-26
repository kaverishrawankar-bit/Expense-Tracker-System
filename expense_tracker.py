
import tkinter as tk
from tkinter import ttk, messagebox
from openpyxl import Workbook, load_workbook
from datetime import datetime
import os


# =========================================================
# EXPENSE TRACKER SYSTEM
# Python + Tkinter + Excel
# =========================================================

EXCEL_FILE = "ExpenseTracker.xlsx"

USERNAME = "admin"
PASSWORD = "admin123"


# =========================================================
# CREATE EXCEL FILE
# =========================================================

def create_excel_file():
    if not os.path.exists(EXCEL_FILE):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Expenses"

        headers = [
            "ID",
            "Date",
            "Category",
            "Description",
            "Amount",
            "Payment Mode"
        ]

        sheet.append(headers)
        workbook.save(EXCEL_FILE)
        workbook.close()


# =========================================================
# GET NEXT ID
# =========================================================

def get_next_id():
    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    ids = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            try:
                ids.append(int(row[0]))
            except:
                pass

    workbook.close()

    if ids:
        return max(ids) + 1

    return 1


# =========================================================
# CLEAR / RESET FORM
# =========================================================

def clear_form():
    date_entry.delete(0, tk.END)
    date_entry.insert(
        0,
        datetime.now().strftime("%d-%m-%Y")
    )

    category_combo.set("")
    description_entry.delete(0, tk.END)
    amount_entry.delete(0, tk.END)
    payment_combo.set("")

    selected_id.set("")


# =========================================================
# ADD EXPENSE
# =========================================================

def add_expense():

    date = date_entry.get().strip()
    category = category_combo.get().strip()
    description = description_entry.get().strip()
    amount = amount_entry.get().strip()
    payment_mode = payment_combo.get().strip()

    if not date or not category or not description or not amount or not payment_mode:
        messagebox.showwarning(
            "Missing Information",
            "Please fill all fields."
        )
        return

    try:
        datetime.strptime(date, "%d-%m-%Y")
    except ValueError:
        messagebox.showerror(
            "Invalid Date",
            "Please enter date in DD-MM-YYYY format."
        )
        return

    try:
        amount_value = float(amount)

        if amount_value <= 0:
            raise ValueError

    except ValueError:
        messagebox.showerror(
            "Invalid Amount",
            "Please enter a valid positive amount."
        )
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    new_id = get_next_id()

    sheet.append([
        new_id,
        date,
        category,
        description,
        amount_value,
        payment_mode
    ])

    workbook.save(EXCEL_FILE)
    workbook.close()

    messagebox.showinfo(
        "Success",
        "Expense added successfully."
    )

    clear_form()
    view_records()


# =========================================================
# VIEW RECORDS
# =========================================================

def view_records():

    for item in tree.get_children():
        tree.delete(item)

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    total = 0
    count = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] is not None:

            tree.insert(
                "",
                tk.END,
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5]
                )
            )

            try:
                total += float(row[4])
            except:
                pass

            count += 1

    workbook.close()

    total_label.config(
        text=f"Total Expenses: ₹{total:.2f}"
    )

    count_label.config(
        text=f"Total Records: {count}"
    )


# =========================================================
# SEARCH EXPENSES
# =========================================================

def search_expenses():

    search_text = search_entry.get().strip().lower()

    if not search_text:
        view_records()
        return

    for item in tree.get_children():
        tree.delete(item)

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    total = 0
    count = 0

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if row[0] is None:
            continue

        values = [
            str(row[0]),
            str(row[1]),
            str(row[2]),
            str(row[3]),
            str(row[4]),
            str(row[5])
        ]

        combined_text = " ".join(values).lower()

        if search_text in combined_text:

            tree.insert(
                "",
                tk.END,
                values=(
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5]
                )
            )

            try:
                total += float(row[4])
            except:
                pass

            count += 1

    workbook.close()

    total_label.config(
        text=f"Search Total: ₹{total:.2f}"
    )

    count_label.config(
        text=f"Records Found: {count}"
    )


# =========================================================
# SELECT RECORD
# =========================================================

def select_record(event=None):

    selected = tree.selection()

    if not selected:
        return

    item = tree.item(selected[0])
    values = item["values"]

    if not values:
        return

    selected_id.set(str(values[0]))

    date_entry.delete(0, tk.END)
    date_entry.insert(0, str(values[1]))

    category_combo.set(str(values[2]))

    description_entry.delete(0, tk.END)
    description_entry.insert(0, str(values[3]))

    amount_entry.delete(0, tk.END)
    amount_entry.insert(0, str(values[4]))

    payment_combo.set(str(values[5]))


# =========================================================
# UPDATE EXPENSE
# =========================================================

def update_expense():

    if not selected_id.get():
        messagebox.showwarning(
            "No Record Selected",
            "Please select a record from the table first."
        )
        return

    date = date_entry.get().strip()
    category = category_combo.get().strip()
    description = description_entry.get().strip()
    amount = amount_entry.get().strip()
    payment_mode = payment_combo.get().strip()

    if not date or not category or not description or not amount or not payment_mode:
        messagebox.showwarning(
            "Missing Information",
            "Please fill all fields."
        )
        return

    try:
        datetime.strptime(date, "%d-%m-%Y")
    except ValueError:
        messagebox.showerror(
            "Invalid Date",
            "Please enter date in DD-MM-YYYY format."
        )
        return

    try:
        amount_value = float(amount)

        if amount_value <= 0:
            raise ValueError

    except ValueError:
        messagebox.showerror(
            "Invalid Amount",
            "Please enter a valid positive amount."
        )
        return

    record_id = int(selected_id.get())

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    found = False

    for row in sheet.iter_rows(min_row=2):

        if row[0].value == record_id:

            row[1].value = date
            row[2].value = category
            row[3].value = description
            row[4].value = amount_value
            row[5].value = payment_mode

            found = True
            break

    workbook.save(EXCEL_FILE)
    workbook.close()

    if found:

        messagebox.showinfo(
            "Success",
            "Expense updated successfully."
        )

        clear_form()
        view_records()

    else:

        messagebox.showerror(
            "Error",
            "Record not found."
        )


# =========================================================
# DELETE EXPENSE
# =========================================================

def delete_expense():

    selected = tree.selection()

    if not selected:

        messagebox.showwarning(
            "No Record Selected",
            "Please select a record to delete."
        )

        return

    item = tree.item(selected[0])
    values = item["values"]

    if not values:
        return

    record_id = int(values[0])

    confirmation = messagebox.askyesno(
        "Confirm Deletion",
        "Are you sure you want to delete this expense?"
    )

    if not confirmation:
        return

    workbook = load_workbook(EXCEL_FILE)
    sheet = workbook.active

    deleted = False

    for row in range(2, sheet.max_row + 1):

        if sheet.cell(row=row, column=1).value == record_id:

            sheet.delete_rows(row, 1)
            deleted = True
            break

    workbook.save(EXCEL_FILE)
    workbook.close()

    if deleted:

        messagebox.showinfo(
            "Deleted",
            "Expense deleted successfully."
        )

        clear_form()
        view_records()

    else:

        messagebox.showerror(
            "Error",
            "Record not found."
        )


# =========================================================
# LOGOUT
# =========================================================

def logout():

    answer = messagebox.askyesno(
        "Logout",
        "Are you sure you want to logout?"
    )

    if answer:

        dashboard_frame.pack_forget()
        login_frame.pack(
            fill="both",
            expand=True
        )

        username_entry.delete(0, tk.END)
        password_entry.delete(0, tk.END)

        username_entry.focus()


# =========================================================
# LOGIN
# =========================================================

def login():

    username = username_entry.get().strip()
    password = password_entry.get().strip()

    if username == USERNAME and password == PASSWORD:

        login_frame.pack_forget()

        dashboard_frame.pack(
            fill="both",
            expand=True
        )

        clear_form()
        view_records()

    else:

        messagebox.showerror(
            "Login Failed",
            "Invalid username or password."
        )


# =========================================================
# MAIN WINDOW
# =========================================================

create_excel_file()

root = tk.Tk()

root.title("Expense Tracker System")
root.geometry("1000x700")
root.minsize(850, 600)

root.configure(bg="#f2f2f2")


# =========================================================
# LOGIN PAGE
# =========================================================

login_frame = tk.Frame(
    root,
    bg="#f2f2f2"
)

login_frame.pack(
    fill="both",
    expand=True
)


login_box = tk.Frame(
    login_frame,
    bg="white",
    padx=40,
    pady=40
)

login_box.place(
    relx=0.5,
    rely=0.5,
    anchor="center"
)


title_label = tk.Label(
    login_box,
    text="Expense Tracker System",
    font=("Arial", 24, "bold"),
    bg="white"
)

title_label.pack(pady=(0, 25))


login_subtitle = tk.Label(
    login_box,
    text="Login",
    font=("Arial", 18, "bold"),
    bg="white"
)

login_subtitle.pack(pady=10)


username_label = tk.Label(
    login_box,
    text="Username",
    font=("Arial", 12),
    bg="white"
)

username_label.pack(anchor="w")


username_entry = tk.Entry(
    login_box,
    font=("Arial", 12),
    width=30
)

username_entry.pack(
    pady=(5, 15)
)


password_label = tk.Label(
    login_box,
    text="Password",
    font=("Arial", 12),
    bg="white"
)

password_label.pack(anchor="w")


password_entry = tk.Entry(
    login_box,
    font=("Arial", 12),
    width=30,
    show="*"
)

password_entry.pack(
    pady=(5, 20)
)


login_button = tk.Button(
    login_box,
    text="LOGIN",
    font=("Arial", 12, "bold"),
    width=20,
    command=login
)

login_button.pack(pady=10)


login_info = tk.Label(
    login_box,
    text="Username: admin   Password: admin123",
    font=("Arial", 9),
    bg="white"
)

login_info.pack(pady=10)


# =========================================================
# DASHBOARD
# =========================================================

dashboard_frame = tk.Frame(
    root,
    bg="#f2f2f2"
)


# =========================================================
# HEADER
# =========================================================

header = tk.Frame(
    dashboard_frame,
    bg="#333333",
    height=70
)

header.pack(
    fill="x"
)

header.pack_propagate(False)

# =========================================================
# NAVIGATION MENU
# =========================================================

nav_frame = tk.Frame(
    dashboard_frame,
    bg="#dddddd",
    height=45
)

nav_frame.pack(
    fill="x"
)

nav_frame.pack_propagate(False)


def go_dashboard():
    view_records()


def go_add_expense():
    clear_form()


def go_view_records():
    view_records()


def go_search():
    search_entry.focus()


tk.Button(
    nav_frame,
    text="Dashboard",
    command=go_dashboard
).pack(
    side="left",
    padx=5,
    pady=5
)


tk.Button(
    nav_frame,
    text="Add Expense",
    command=go_add_expense
).pack(
    side="left",
    padx=5,
    pady=5
)


tk.Button(
    nav_frame,
    text="View Records",
    command=go_view_records
).pack(
    side="left",
    padx=5,
    pady=5
)


tk.Button(
    nav_frame,
    text="Search",
    command=go_search
).pack(
    side="left",
    padx=5,
    pady=5
)

dashboard_title = tk.Label(
    header,
    text="Expense Tracker System",
    font=("Arial", 22, "bold"),
    fg="white",
    bg="#333333"
)

dashboard_title.pack(
    side="left",
    padx=20
)


logout_button = tk.Button(
    header,
    text="Logout",
    font=("Arial", 11, "bold"),
    command=logout
)

logout_button.pack(
    side="right",
    padx=20
)


# =========================================================
# MAIN AREA
# =========================================================

main_area = tk.Frame(
    dashboard_frame,
    bg="#f2f2f2"
)

main_area.pack(
    fill="both",
    expand=True,
    padx=15,
    pady=15
)


# =========================================================
# EXPENSE FORM
# =========================================================

form_frame = tk.LabelFrame(
    main_area,
    text="Expense Details",
    font=("Arial", 12, "bold"),
    bg="#f2f2f2",
    padx=15,
    pady=15
)

form_frame.pack(
    side="left",
    fill="y",
    padx=(0, 10)
)


selected_id = tk.StringVar()


# Date
tk.Label(
    form_frame,
    text="Date (DD-MM-YYYY)",
    bg="#f2f2f2",
    font=("Arial", 10)
).pack(anchor="w")

date_entry = tk.Entry(
    form_frame,
    width=25,
    font=("Arial", 11)
)

date_entry.pack(
    pady=(3, 12)
)


# Category
tk.Label(
    form_frame,
    text="Category",
    bg="#f2f2f2",
    font=("Arial", 10)
).pack(anchor="w")

category_combo = ttk.Combobox(
    form_frame,
    width=22,
    values=[
        "Food",
        "Travel",
        "Shopping",
        "Education",
        "Entertainment",
        "Bills",
        "Health",
        "Other"
    ],
    state="readonly"
)

category_combo.pack(
    pady=(3, 12)
)


# Description
tk.Label(
    form_frame,
    text="Description",
    bg="#f2f2f2",
    font=("Arial", 10)
).pack(anchor="w")

description_entry = tk.Entry(
    form_frame,
    width=25,
    font=("Arial", 11)
)

description_entry.pack(
    pady=(3, 12)
)


# Amount
tk.Label(
    form_frame,
    text="Amount (₹)",
    bg="#f2f2f2",
    font=("Arial", 10)
).pack(anchor="w")

amount_entry = tk.Entry(
    form_frame,
    width=25,
    font=("Arial", 11)
)

amount_entry.pack(
    pady=(3, 12)
)


# Payment Mode
tk.Label(
    form_frame,
    text="Payment Mode",
    bg="#f2f2f2",
    font=("Arial", 10)
).pack(anchor="w")

payment_combo = ttk.Combobox(
    form_frame,
    width=22,
    values=[
        "Cash",
        "UPI",
        "Debit Card",
        "Credit Card",
        "Net Banking",
        "Other"
    ],
    state="readonly"
)

payment_combo.pack(
    pady=(3, 15)
)


# =========================================================
# FORM BUTTONS
# =========================================================

add_button = tk.Button(
    form_frame,
    text="Add Expense",
    width=22,
    command=add_expense
)

add_button.pack(pady=5)


update_button = tk.Button(
    form_frame,
    text="Update Expense",
    width=22,
    command=update_expense
)

update_button.pack(pady=5)


delete_button = tk.Button(
    form_frame,
    text="Delete Expense",
    width=22,
    command=delete_expense
)

delete_button.pack(pady=5)


clear_button = tk.Button(
    form_frame,
    text="Clear / Reset",
    width=22,
    command=clear_form
)

clear_button.pack(pady=5)


# =========================================================
# RIGHT SIDE
# =========================================================

right_frame = tk.Frame(
    main_area,
    bg="#f2f2f2"
)

right_frame.pack(
    side="left",
    fill="both",
    expand=True
)


# =========================================================
# SEARCH
# =========================================================

search_frame = tk.Frame(
    right_frame,
    bg="#f2f2f2"
)

search_frame.pack(
    fill="x",
    pady=(0, 10)
)


tk.Label(
    search_frame,
    text="Search:",
    font=("Arial", 11, "bold"),
    bg="#f2f2f2"
).pack(side="left")


search_entry = tk.Entry(
    search_frame,
    font=("Arial", 11),
    width=30
)

search_entry.pack(
    side="left",
    padx=8
)


search_button = tk.Button(
    search_frame,
    text="Search",
    command=search_expenses
)

search_button.pack(
    side="left",
    padx=3
)


show_all_button = tk.Button(
    search_frame,
    text="Show All",
    command=view_records
)

show_all_button.pack(
    side="left",
    padx=3
)


# =========================================================
# SUMMARY
# =========================================================

summary_frame = tk.Frame(
    right_frame,
    bg="#f2f2f2"
)

summary_frame.pack(
    fill="x",
    pady=(0, 10)
)


total_label = tk.Label(
    summary_frame,
    text="Total Expenses: ₹0.00",
    font=("Arial", 12, "bold"),
    bg="#f2f2f2"
)

total_label.pack(
    side="left",
    padx=10
)


count_label = tk.Label(
    summary_frame,
    text="Total Records: 0",
    font=("Arial", 12, "bold"),
    bg="#f2f2f2"
)

count_label.pack(
    side="right",
    padx=10
)


# =========================================================
# TABLE
# =========================================================

table_frame = tk.Frame(
    right_frame
)

table_frame.pack(
    fill="both",
    expand=True
)


columns = (
    "ID",
    "Date",
    "Category",
    "Description",
    "Amount",
    "Payment Mode"
)


tree = ttk.Treeview(
    table_frame,
    columns=columns,
    show="headings"
)


for column in columns:

    tree.heading(
        column,
        text=column
    )


tree.column(
    "ID",
    width=50,
    anchor="center"
)

tree.column(
    "Date",
    width=100,
    anchor="center"
)

tree.column(
    "Category",
    width=110,
    anchor="center"
)

tree.column(
    "Description",
    width=180,
    anchor="center"
)

tree.column(
    "Amount",
    width=100,
    anchor="center"
)

tree.column(
    "Payment Mode",
    width=120,
    anchor="center"
)


# =========================================================
# SCROLLBAR
# =========================================================

scrollbar = ttk.Scrollbar(
    table_frame,
    orient="vertical",
    command=tree.yview
)

tree.configure(
    yscrollcommand=scrollbar.set
)


tree.pack(
    side="left",
    fill="both",
    expand=True
)

scrollbar.pack(
    side="right",
    fill="y"
)


# Select a record
tree.bind(
    "<ButtonRelease-1>",
    select_record
)


# =========================================================
# START PROGRAM
# =========================================================

username_entry.focus()

root.mainloop()